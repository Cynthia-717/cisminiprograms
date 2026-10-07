import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from scripts.import_pdfs import extract_courses_from_excel, extract_requirements_from_excel, infer_domain, parse_requirement_text


class ExcelImportTests(unittest.TestCase):
    def test_extract_courses_from_excel_second_sheet(self):
        wb = Workbook()
        ws1 = wb.active
        ws1.title = "第一張工作表"
        ws1["A1"] = "ignore"

        ws2 = wb.create_sheet("課程規劃表")
        ws2.append(["類別", "科目名稱", "學分"])
        ws2.append(["基礎", "資料結構", 2])
        ws2.append(["核心", "機器學習", 3])
        ws2.append(["應用", "AI 專題", 2])

        tmp_dir = Path(tempfile.gettempdir()) / "microprogram-test"
        tmp_dir.mkdir(exist_ok=True)
        path = tmp_dir / "人工智慧微學程 1142.xlsx"
        wb.save(path)

        courses = extract_courses_from_excel(path)

        self.assertEqual(
            [course["name"] for course in courses],
            ["資料結構", "機器學習", "AI專題"],
        )
        self.assertEqual(
            [course["category"] for course in courses],
            ["基礎", "核心", "應用"],
        )

    def test_extracts_program_courses_not_offering_units(self):
        wb = Workbook()
        wb.active.title = "封面"
        ws = wb.create_sheet("課程規劃表")
        ws.append(["領域名稱", "領域應修學分數", "開課單位", "學程課程", "課號", "選別", "科目學分"])
        ws.append(["基礎", 2, "藥學系學士班共同整合課程", "基礎毒理學", "Z0300028", "選修", 3])
        ws.append(["基礎", 2, "醫學科學研究所博士班", "基礎毒理學", "3195E008", "必修", 3])
        ws.append(["核心", 2, "跨領域學院學士班課程", "臨床試驗實務概論", "XB500169", "選修", 2])
        ws.append(["應用", 2, "不應成為課名", "", "", "選修", 2])

        tmp_dir = Path(tempfile.gettempdir()) / "microprogram-test"
        tmp_dir.mkdir(exist_ok=True)
        path = tmp_dir / "臨床試驗微學程 1151.xlsx"
        wb.save(path)

        courses = extract_courses_from_excel(path)

        self.assertEqual(
            courses,
            [
                {"category": "基礎", "name": "基礎毒理學", "credit": 3},
                {"category": "核心", "name": "臨床試驗實務概論", "credit": 2},
            ],
        )

    def test_extracts_requirements_from_first_sheet(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["適用學年期", "學程名稱", "修業規定", "至少需修畢學分數"])
        ws.append(["1151", "測試微學程", "本微學程需修習6學分，基礎、核心、應用各需修習2學分。", 6])
        wb.create_sheet("課程規劃表")

        tmp_dir = Path(tempfile.gettempdir()) / "microprogram-test"
        tmp_dir.mkdir(exist_ok=True)
        path = tmp_dir / "測試微學程 1151.xlsx"
        wb.save(path)

        requirements = extract_requirements_from_excel(path)

        self.assertEqual(requirements["totalCredits"], 6)
        self.assertEqual(requirements["perCategoryCredits"], {"基礎": 2, "核心": 2, "應用": 2})
        self.assertEqual(requirements["note"], "本微學程需修習6學分，基礎、核心、應用各需修習2學分。")

    def test_extracts_required_category_courses_and_total_from_prose(self):
        requirements = parse_requirement_text(
            "本學程需修習8學分，基礎及核心課程至少各需選修1門，應用課程至少需修習2學分。"
        )

        self.assertEqual(requirements["totalCredits"], 8)
        self.assertEqual(requirements["minCoursesPerCategory"], {"基礎": 1, "核心": 1, "應用": 0})
        self.assertEqual(requirements["perCategoryCredits"], {"基礎": None, "核心": None, "應用": 2})

    def test_extracts_chinese_course_counts_and_required_category_without_amount(self):
        requirements = parse_requirement_text(
            "本學程需修習7學分，基礎必修兩門；應用課程為取得微學程之必要項目。"
        )

        self.assertEqual(requirements["minCoursesPerCategory"]["基礎"], 2)
        self.assertTrue(requirements["requiredCategories"]["基礎"])
        self.assertTrue(requirements["requiredCategories"]["應用"])
        self.assertIsNone(requirements["perCategoryCredits"]["應用"])

    def test_does_not_apply_one_category_count_to_prior_category_list(self):
        requirements = parse_requirement_text(
            "本學程需修習7學分，包含基礎、核心、應用課程，基礎必修兩門。"
        )

        self.assertEqual(requirements["minCoursesPerCategory"], {"基礎": 2, "核心": 0, "應用": 0})
        self.assertEqual(requirements["perCategoryCredits"], {"基礎": None, "核心": None, "應用": None})

    def test_marks_only_the_category_explicitly_called_necessary(self):
        requirements = parse_requirement_text(
            "本學程需修習6學分，包含基礎、核心、應用課程，其中應用課程為取得微學程之必要項目。"
        )

        self.assertEqual(requirements["requiredCategories"], {"基礎": False, "核心": False, "應用": True})

    def test_does_not_infer_category_minimum_when_rule_only_lists_categories(self):
        requirements = parse_requirement_text("本學程需修習8學分，包含基礎、核心、應用課程，其餘學分可自行選修。")

        self.assertEqual(requirements["totalCredits"], 8)
        self.assertEqual(requirements["minCoursesPerCategory"], {"基礎": 0, "核心": 0, "應用": 0})
        self.assertEqual(requirements["perCategoryCredits"], {"基礎": None, "核心": None, "應用": None})

    def test_infer_domain_from_microprogram_folder_name(self):
        self.assertEqual(
            infer_domain(Path("excels/微學程清單/自主學習微學程/自主學習微學程1151.xls")),
            "未分類",
        )
        self.assertEqual(
            infer_domain(Path("excels/微學程清單/人工智慧微學程/人工智慧微學程1151.xls")),
            "人工智慧領域",
        )
        self.assertEqual(
            infer_domain(Path("excels/微學程清單/氣候變遷、綠能永續與健康微學程/氣候變遷、綠能永續與健康微學程1151.xls")),
            "永續發展領域",
        )


if __name__ == "__main__":
    unittest.main()
